import os
from typing import Any, Optional
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from app.tools.base import Tool, ToolResult


class DocumentTool(Tool):
    name: str = "create_document"
    description: str = (
        "Creates a Word document (.docx) or PDF from provided content. "
        "Use this when the user asks to generate, write, or create a document, report, letter, or file."
    )

    def run(
        self,
        content: Optional[str] = None,
        format: Optional[str] = None,
        filename: Optional[str] = None,
        **kwargs: Any,
    ) -> ToolResult:
        """Generates a .docx or .pdf document from markdown-like structured text."""
        doc_content = content if content is not None else kwargs.get("content", "")
        doc_format = (format if format is not None else (kwargs.get("format") or "docx")).lower().strip()
        raw_filename = filename if filename is not None else (kwargs.get("filename") or "document")

        if not doc_content or not isinstance(doc_content, str) or not doc_content.strip():
            return ToolResult(
                success=False,
                output=None,
                error="No document content provided.",
            )

        try:
            base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
            output_dir = os.path.join(base_dir, "generated_docs")
            os.makedirs(output_dir, exist_ok=True)

            clean_filename = os.path.splitext(os.path.basename(str(raw_filename).strip() or "document"))[0]

            if doc_format in ("docx", ".docx", "word"):
                save_path = os.path.join(output_dir, f"{clean_filename}.docx")
                doc = docx.Document()

                for raw_line in doc_content.split("\n"):
                    line = raw_line.strip()
                    if not line:
                        continue
                    if line.startswith("### "):
                        doc.add_heading(line[4:].strip(), level=2)
                    elif line.startswith("## "):
                        doc.add_heading(line[3:].strip(), level=1)
                    elif line.startswith("- "):
                        doc.add_paragraph(line[2:].strip(), style="List Bullet")
                    else:
                        doc.add_paragraph(line)

                doc.save(save_path)
                return ToolResult(
                    success=True,
                    output=save_path,
                    error=None,
                )

            elif doc_format in ("pdf", ".pdf"):
                save_path = os.path.join(output_dir, f"{clean_filename}.pdf")
                doc_template = SimpleDocTemplate(save_path, pagesize=letter)
                styles = getSampleStyleSheet()
                story = []

                for raw_line in doc_content.split("\n"):
                    line = raw_line.strip()
                    if not line:
                        continue
                    if line.startswith("### "):
                        story.append(Paragraph(line[4:].strip(), styles["Heading2"]))
                        story.append(Spacer(1, 6))
                    elif line.startswith("## "):
                        story.append(Paragraph(line[3:].strip(), styles["Heading1"]))
                        story.append(Spacer(1, 8))
                    elif line.startswith("- "):
                        bullet_text = f"\u2022 {line[2:].strip()}"
                        story.append(Paragraph(bullet_text, styles["Normal"]))
                        story.append(Spacer(1, 4))
                    else:
                        story.append(Paragraph(line, styles["Normal"]))
                        story.append(Spacer(1, 6))

                doc_template.build(story)
                return ToolResult(
                    success=True,
                    output=save_path,
                    error=None,
                )

            else:
                return ToolResult(
                    success=False,
                    output=None,
                    error=f"Unsupported format '{doc_format}'. Supported formats are 'docx' and 'pdf'.",
                )

        except Exception as e:
            return ToolResult(
                success=False,
                output=None,
                error=f"Document generation failed: {str(e)}",
            )
